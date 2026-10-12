"""Bank file parsing, deliberately independent of Odoo and accounting writes."""
import csv
import io
import math
import re
import unicodedata
import zipfile
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal, InvalidOperation


MAX_FILE_BYTES = 10 * 1024 * 1024
MAX_XLSX_BYTES = 50 * 1024 * 1024
MAX_ROWS = 10000
MAX_COLUMNS = 100
REQUIRED = {'FECHA', 'MONEDA', 'MOVIMIENTO'}
MONEY = re.compile(r'^[+-]?(?:\d+|\d{1,3}(?:,\d{3})+)(?:\.\d+)?$')


class FileError(ValueError):
    pass


@dataclass(frozen=True)
class Row:
    number: int
    movement: str
    date: date
    amount: Decimal
    reference: str
    label: str


@dataclass(frozen=True)
class Issue:
    number: int
    movement: str
    message: str


@dataclass
class Parsed:
    rows: list
    issues: list
    read_count: int


def text(value):
    return '' if value is None or value is False else str(value).strip()


def normalize_currency(value):
    value = text(value).upper()
    return 'MXN' if value == 'MXP' else value


def header(value):
    value = unicodedata.normalize('NFKD', text(value).upper())
    return ' '.join(''.join(c for c in value if not unicodedata.combining(c)).split())


def parse_amount(value, *, zero_if_empty=False):
    if isinstance(value, bool):
        raise ValueError('Importe inválido.')
    if zero_if_empty and (value is None or text(value) in ('', '-')):
        return Decimal(0)
    if isinstance(value, (int, float, Decimal)):
        result = Decimal(str(value))
    else:
        raw = text(value)
        negative = raw.startswith('(') and raw.endswith(')')
        if negative:
            raw = raw[1:-1].strip()
        raw = raw.removeprefix('$').strip()
        if not MONEY.fullmatch(raw) or (negative and raw.startswith(('+', '-'))):
            raise ValueError('Importe inválido; usa, por ejemplo, $18,872.11.')
        try:
            result = Decimal(raw.replace(',', ''))
        except InvalidOperation as exc:
            raise ValueError('Importe inválido.') from exc
        if negative:
            result = -result
    if not result.is_finite() or not math.isfinite(float(result)):
        raise ValueError('El importe debe ser un número finito.')
    return result


def parse_date(value):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    raw = text(value)
    if not re.fullmatch(r'\d{2}/\d{2}/\d{4}', raw):
        raise ValueError('FECHA debe tener formato DD/MM/AAAA o ser una fecha de Excel.')
    try:
        return datetime.strptime(raw, '%d/%m/%Y').date()
    except ValueError as exc:
        raise ValueError('FECHA no es una fecha válida (DD/MM/AAAA).') from exc


def _identifier(cell, name):
    """Preserve text and explicit Excel zero padding; never guess lost zeroes."""
    value = getattr(cell, 'value', cell)
    if getattr(cell, 'data_type', None) in ('f', 'e'):
        raise ValueError(f'{name}: no se admiten fórmulas ni errores de Excel.')
    if isinstance(value, bool):
        raise ValueError(f'{name} debe ser texto.')
    if isinstance(value, (int, float)):
        if not math.isfinite(value) or value != int(value) or value < 0 or value >= 10**15:
            raise ValueError(f'{name}: guarda el identificador como texto para conservar su precisión.')
        number_format = getattr(cell, 'number_format', '')
        if re.fullmatch('0+', number_format):
            return str(int(value)).zfill(len(number_format))
        return str(int(value))
    return text(value)


def _value(cell):
    if getattr(cell, 'data_type', None) in ('f', 'e'):
        raise ValueError('No se admiten fórmulas ni errores de Excel en los campos utilizados.')
    return getattr(cell, 'value', cell)


def _csv_rows(data):
    try:
        source = data.decode('utf-8-sig')
    except UnicodeDecodeError:
        try:
            source = data.decode('cp1252')
        except UnicodeDecodeError as exc:
            raise FileError('El CSV debe estar codificado como UTF-8 o CP1252.') from exc
    try:
        dialect = csv.Sniffer().sniff(source[:65536], delimiters=',;\t')
    except csv.Error:
        dialect = csv.excel
    try:
        reader = csv.reader(io.StringIO(source, newline=''), dialect, strict=True)
        for row in reader:
            yield reader.line_num, row
    except csv.Error as exc:
        raise FileError(f'CSV inválido cerca de la línea {reader.line_num}: {exc}') from exc


def _xlsx_rows(data):
    from openpyxl import load_workbook

    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            if sum(item.file_size for item in archive.infolist()) > MAX_XLSX_BYTES:
                raise FileError('El XLSX descomprimido supera 50 MiB.')
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=False)
    except (OSError, ValueError, KeyError, TypeError, IndexError, SyntaxError, zipfile.BadZipFile) as exc:
        raise FileError(f'No se pudo leer el XLSX: {exc}') from exc
    try:
        # A single nonempty worksheet is intentional; silently dropping others
        # would hide transactions from validation.
        selected = None
        for sheet in workbook.worksheets:
            if sheet.max_column and sheet.max_column > MAX_COLUMNS:
                raise FileError('El archivo supera 100 columnas.')
            for number, row in enumerate(sheet.iter_rows(), 1):
                if number > MAX_ROWS + 1:
                    raise FileError('El archivo supera 10,000 filas por hoja.')
                if any(text(cell.value) for cell in row):
                    if selected is not None and selected != sheet.title:
                        raise FileError('El XLSX debe contener una sola hoja con datos.')
                    selected = sheet.title
                    yield number, row
    except FileError:
        raise
    except (OSError, ValueError, KeyError, TypeError, IndexError, SyntaxError, zipfile.BadZipFile) as exc:
        raise FileError(f'No se pudo leer el contenido del XLSX: {exc}') from exc
    finally:
        workbook.close()


def parse_file(data, filename, *, currency):
    if not data:
        raise FileError('El archivo está vacío.')
    if len(data) > MAX_FILE_BYTES:
        raise FileError('El archivo supera 10 MiB.')
    extension = text(filename).lower().rsplit('.', 1)[-1]
    if extension not in ('csv', 'xlsx'):
        raise FileError('Selecciona un archivo CSV o XLSX.')
    source = iter(_csv_rows(data) if extension == 'csv' else _xlsx_rows(data))
    rows, issues = [], []
    columns = None
    read_count = 0
    for number, cells in source:
        if len(cells) > MAX_COLUMNS:
            raise FileError('El archivo supera 100 columnas.')
        if not any(text(getattr(cell, 'value', cell)) for cell in cells):
            continue
        if columns is None:
            try:
                names = [header(_value(cell)) for cell in cells]
            except ValueError as exc:
                raise FileError(f'Encabezado inválido: {exc}') from exc
            nonempty = [name for name in names if name]
            if len(set(nonempty)) != len(nonempty):
                raise FileError('Hay encabezados duplicados.')
            columns = {name: i for i, name in enumerate(names) if name}
            missing = REQUIRED - columns.keys()
            if 'IMPORTE' not in columns:
                missing |= {'DEPOSITOS', 'RETIROS'} - columns.keys()
            if missing:
                raise FileError('Faltan encabezados obligatorios: ' + ', '.join(sorted(missing)))
            continue
        read_count += 1
        if read_count > MAX_ROWS:
            raise FileError('El archivo supera 10,000 movimientos.')

        def cell(name):
            i = columns.get(name)
            return cells[i] if i is not None and i < len(cells) else None

        movement = ''
        try:
            if any(text(getattr(c, 'value', c)) for c in cells[len(names):]):
                raise ValueError('Hay valores sin encabezado; revisa el separador CSV.')
            movement = _identifier(cell('MOVIMIENTO'), 'MOVIMIENTO')
            if not movement:
                raise ValueError('MOVIMIENTO es obligatorio.')
            if normalize_currency(_value(cell('MONEDA'))) != normalize_currency(currency):
                raise ValueError('MONEDA no corresponde a la moneda del diario.')
            row_date = parse_date(_value(cell('FECHA')))
            # Use the bank's original columns even if a manually added IMPORTE
            # is present. Credits are positive and debits are negative.
            if {'DEPOSITOS', 'RETIROS'} <= columns.keys():
                deposit = parse_amount(_value(cell('DEPOSITOS')), zero_if_empty=True)
                withdrawal = parse_amount(_value(cell('RETIROS')), zero_if_empty=True)
                if deposit < 0 or withdrawal < 0:
                    raise ValueError('DEPÓSITOS y RETIROS deben ser magnitudes positivas o cero.')
                amount = deposit - withdrawal
            else:
                amount = parse_amount(_value(cell('IMPORTE')))
            reference = text(_value(cell('REFERENCIA')))
            label = ' / '.join(filter(None, (
                text(_value(cell('DESCRIPCION'))),
                text(_value(cell('DESCRIPCION DETALLADA'))),
            ))) or reference or movement
            rows.append(Row(number, movement, row_date, amount, reference, label))
        except (ValueError, InvalidOperation) as exc:
            issues.append(Issue(number, movement, str(exc)))
    if columns is None or not read_count:
        raise FileError('El archivo no contiene movimientos.')
    return Parsed(rows, issues, read_count)
